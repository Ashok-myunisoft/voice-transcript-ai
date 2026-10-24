from pathlib import Path
from docx import Document
from docx.shared import Pt


def _speaker_label(speaker_id, speakers):
    speaker = next((s for s in speakers if s["speaker_id"] == speaker_id), None)
    return speaker["name"] if speaker and speaker.get("name") else (speaker_id or "Unknown speaker")


def generate_reports(directory: Path, meeting_id: str, transcript: list[dict], speakers: list[dict], summary: dict,
                     actions: list[dict], duration: float) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    txt = ["MEETING TRANSCRIPT", f"Meeting ID: {meeting_id}", f"Duration: {duration:.0f} seconds", ""]
    for seg in transcript:
        txt.extend([f"[{seg['start']:02.0f}s - {seg['end']:02.0f}s] {_speaker_label(seg.get('speaker_id'), speakers)}:", seg["text"], ""])
    (directory / "transcript.txt").write_text("\n".join(txt), encoding="utf-8")

    doc = Document()
    doc.add_heading("Meeting Report", 0)
    doc.add_paragraph(f"Meeting ID: {meeting_id} | Duration: {duration:.0f} seconds")
    doc.add_heading("Participants", level=1)
    for speaker in speakers:
        doc.add_paragraph(f"{speaker.get('name') or speaker['speaker_id']} ({speaker['identity_status']})", style="List Bullet")
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(summary["overview"])
    doc.add_heading("Key Discussion Points", level=1)
    for point in summary["key_discussion_points"]:
        doc.add_paragraph(point, style="List Bullet")
    doc.add_heading("Decisions", level=1)
    if summary["decisions"]:
        for decision in summary["decisions"]:
            doc.add_paragraph(str(decision), style="List Bullet")
    else:
        doc.add_paragraph("No decisions were confidently extracted.")
    doc.add_heading("Action Items", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ["Task", "Owner", "Deadline", "Status"]):
        cell.text = value
    for item in actions:
        cells = table.add_row().cells
        for cell, value in zip(cells, [item["task"], item.get("owner_name") or item.get("owner_speaker_id") or "Unknown",
                                       item.get("deadline") or "Unknown", item.get("status") or "Unknown"]):
            cell.text = str(value)
    doc.add_heading("Pending Questions and Follow-ups", level=1)
    pending = summary.get("pending_questions", [])
    for item in pending:
        doc.add_paragraph(str(item), style="List Bullet")
    if not pending:
        doc.add_paragraph("None extracted.")
    doc.save(directory / "meeting_report.docx")
